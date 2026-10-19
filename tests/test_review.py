import csv
import json

import pytest

from nrag.extract import review
from test_extract import good_extraction


def _docs(n_per=3):
    docs = []
    for design in ("rct", "meta_analysis", "prospective_cohort"):
        for topic in ("weight_management", "maternal_nutrition"):
            for i in range(n_per):
                docs.append({"doc_id": f"PMID:{design[:3]}{topic[:3]}{i}", "pmid": "1", "design": design,
                             "topics": [topic], "title": "t", "abstract": "a", "mesh": ["Adult", "Pregnancy"],
                             "sample_size": 100})
    return docs


def test_stratified_sample_balanced_deterministic_and_nested():
    docs = _docs()
    s12 = review.stratified_sample(docs, 12, seed=1)
    strata = {(d["design"], d["topics"][0]) for d in s12[:6]}
    assert len(strata) == 6                                   # first round covers every stratum
    assert [d["doc_id"] for d in review.stratified_sample(docs, 12, seed=1)] == [d["doc_id"] for d in s12]
    assert [d["doc_id"] for d in review.stratified_sample(docs, 6, seed=1)] == [d["doc_id"] for d in s12[:6]]
    assert len(review.stratified_sample(docs, 500, seed=1)) == len(docs)


def test_render_and_get_field():
    e = good_extraction()
    assert review.render(review.get_field(e, "population.tags"), "population.tags") == "indian; prediabetes; south_asian"
    assert review.render(review.get_field(e, "outcomes"), "outcomes") == "Diabetes incidence [hard]"
    assert review.render(review.get_field(e, "risk_of_bias_signals.pre_registered"), "x") == "yes"
    assert review.render(None, "intervention") == "null"
    assert review.render(52.0, "duration_weeks") == "52"
    assert review.get_field(None, "a.b") is None


@pytest.mark.parametrize("field, value, gold, ok", [
    ("sample_size", "240", "240", True),
    ("sample_size", "245", "240", True),        # within 5%
    ("sample_size", "300", "240", False),
    ("sample_size", "null", "null", True),
    ("sample_size", "240", "null", False),
    ("population.tags", "indian; pregnant", "pregnant; indian", True),
    ("population.tags", "pregnant", "pregnant; indian", False),
    ("effect_direction", "Benefit", "benefit", True),
])
def test_values_match(field, value, gold, ok):
    assert review.values_match(field, value, gold) == ok


def test_rules_baseline(cfg):
    mesh_tags = cfg["extract"]["rules_baseline_mesh_tags"]
    b = review.rules_baseline({"mesh": ["Adult", "Pregnancy", "India", "Diabetes, Gestational"], "design": "cohort",
                               "sample_size": None}, mesh_tags)
    assert b == {"population.tags": "gestational_diabetes; indian; pregnant; south_asian",
                 "sample_size": "null", "study_design": "retrospective_cohort"}
    assert review.rules_baseline({"mesh": ["Adult", "Humans"], "design": "rct", "sample_size": 80}, mesh_tags)[
        "population.tags"] == "general_adults"


def _write_and_fill(tmp_path, docs, sys_a, sys_b, fill):
    csv_path, key_path = tmp_path / "review.csv", tmp_path / "key.json"
    review.write_review_csv(csv_path, key_path, docs, {"m1": sys_a, "m2": sys_b}, seed=3)
    key = json.loads(key_path.read_text())
    rows = review.read_review(csv_path)
    for row in rows:
        fill(row, key["assignment"][row["doc_id"]])
    return rows, key


def test_review_csv_is_blinded_and_complete(tmp_path):
    docs = _docs(1)
    ext = {d["doc_id"]: good_extraction() for d in docs}
    review.write_review_csv(tmp_path / "r.csv", tmp_path / "k.json", docs, {"m1": ext, "m2": ext}, seed=3)
    rows = list(csv.DictReader((tmp_path / "r.csv").open()))
    assert len(rows) == len(docs) * len(review.REVIEW_FIELDS)
    assert rows[0]["abstract"] == "a" and rows[1]["abstract"] == ""      # abstract only on a doc's first row
    key = json.loads((tmp_path / "k.json").read_text())
    assert {v["A"] for v in key["assignment"].values()} == {"m1", "m2"}  # randomised per document


def test_score_review_end_to_end(tmp_path, cfg):
    docs = _docs(2)      # 12 docs
    good = {d["doc_id"]: good_extraction() for d in docs}
    bad = {d["doc_id"]: good_extraction(effect_direction="no_effect", sample_size=999,
                                        risk_of_bias_signals={"randomisation_described": "yes", "blinding": "not_reported",
                                                              "low_dropout": "not_reported", "pre_registered": "yes"})
           for d in docs}

    def fill(row, assignment):
        # The reviewer agrees with m1 everywhere (m1 is the "truth" here).
        letter_m1 = next(k for k, v in assignment.items() if v == "m1")
        row["verdict"] = "both" if row["A"] == row["B"] else letter_m1

    rows, key = _write_and_fill(tmp_path, docs, good, bad, fill)
    speeds = {"m1": {"median": 30.0}, "m2": {"median": 10.0}}
    baseline = {d["doc_id"]: review.rules_baseline(d, cfg["extract"]["rules_baseline_mesh_tags"]) for d in docs}
    res = review.score_review(rows, key, baseline, speeds, cfg["extract"]["decision"], iters=500, seed=0)

    m1, m2 = res["systems"]["m1"], res["systems"]["m2"]
    assert res["labelled_documents"] == 12 and res["unlabelled_rows"] == 0
    assert m1["fields"]["effect_direction"]["accuracy"] == 1.0 and m2["fields"]["effect_direction"]["accuracy"] == 0.0
    assert m2["fields"]["intervention"]["accuracy"] == 1.0          # identical -> "both"
    assert m2["hallucinated_yes_rate"] == pytest.approx(1 / 3, abs=1e-3)     # randomisation 'yes' when gold not_reported
    assert m1["tags_micro"]["f1"] == 1.0
    assert res["systems"]["rules_baseline"]["fields"]["sample_size"]["accuracy"] == 0.0   # 100 vs gold 240
    c = res["comparison"]
    assert c["critical_accuracy_diff_a_minus_b"] == 0.5               # m1 wins 2 of 4 critical fields
    assert c["ci95"][0] > 0
    assert c["decision"]["choice"] == "m1"                            # faster m2 is credibly worse
    assert c["mcnemar"]["sample_size"] == {"a_only": 12, "b_only": 0, "p": pytest.approx(0.0005, abs=1e-4)}
    assert "Decision" in review.render_report(res)


def test_unlabelled_rows_are_skipped_and_neither_uses_correct_value(tmp_path, cfg):
    docs = _docs(1)[:2]
    ext = {d["doc_id"]: good_extraction() for d in docs}

    def fill(row, assignment):
        if row["field"] == "sample_size":
            row["verdict"], row["correct_value"] = "neither", "300"
        elif row["field"] == "effect_direction":
            row["verdict"] = "both"

    rows, key = _write_and_fill(tmp_path, docs, ext, ext, fill)
    res = review.score_review(rows, key, {}, {}, cfg["extract"]["decision"], iters=100)
    assert res["unlabelled_rows"] == len(rows) - 4
    assert res["systems"]["m1"]["fields"]["sample_size"]["accuracy"] == 0.0     # 240 vs 300 is > 5% off
    assert res["comparison"]["decision"]["choice"] is None                     # no speeds measured


def test_stats_helpers():
    assert review.cohen_kappa([("a", "a"), ("b", "b"), ("a", "b"), ("b", "a")]) == 0.0
    assert review.cohen_kappa([("a", "a"), ("b", "b")]) == 1.0
    assert review.mcnemar_p(0, 0) == 1.0
    assert review.mcnemar_p(10, 0) == pytest.approx(2 / 1024)
    point, lo, hi = review.bootstrap_diff({str(i): (1.0, 0.0) for i in range(10)}, iters=200, seed=0)
    assert (point, lo, hi) == (1.0, 1.0, 1.0)


@pytest.mark.parametrize("diff, lo, hi, choice", [
    (0.10, 0.02, 0.18, "slow"),      # a=slow is credibly better by > 5 points -> slow
    (0.10, -0.02, 0.22, "fast"),     # CI includes zero -> fast
    (0.03, 0.01, 0.05, "fast"),      # better but by < 5 points -> fast
    (-0.10, -0.2, -0.01, "slow"),    # fast model credibly better -> fast... handled below
])
def test_decide(diff, lo, hi, choice):
    speeds = {"slow": {"median": 40}, "fast": {"median": 10}}
    got = review.decide("slow", "fast", diff, lo, hi, speeds, {"max_accuracy_drop": 0.05})["choice"]
    expected = "fast" if diff < 0 else choice
    assert got == expected


def test_add_paper_numbers_and_xlsx(tmp_path):
    pytest.importorskip("openpyxl")
    p = tmp_path / "r.csv"
    p.write_text("doc_id,pmid,field,A,B,verdict,correct_value,notes,title,abstract\n"
                 "PMID:1,1,study_design,rct,rct,both,,,t,a\nPMID:1,1,intervention,x,y,,,,,\n"
                 "PMID:2,2,study_design,rct,rct,,,,t2,a2\n")
    assert review.add_paper_numbers(p) == 2
    rows = review.read_review(p)
    assert [r["paper_no"] for r in rows] == ["1", "1", "2"] and rows[0]["verdict"] == "both"
    review.write_review_xlsx(p, tmp_path / "r.xlsx")
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "r.xlsx").active
    assert ws["A2"].value == 1 and ws["A2"].fill.fgColor.rgb[-6:] != ws["A4"].fill.fgColor.rgb[-6:]
